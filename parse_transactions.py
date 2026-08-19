"""
Generic Bank SMS Transaction Parser.
Supports HDFC, ICICI, SBI, Standard Chartered, Lloyds, and any other bank worldwide.
Parses SMS text files and exports structured transactions to Excel.
Multi-currency: INR, GBP, USD, EUR, AED, etc.
"""
import os
import re
import sys
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load .env for default currency config
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(BASE_DIR, ".env"))

# ── Constants ─────────────────────────────────────────────────────────
DEFAULT_CURRENCY = os.getenv("DEFAULT_CURRENCY", "INR")

MONTH_MAP = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}

# Currency symbols and codes
CURRENCY_SYMBOLS = {
    "₹": "INR", "Rs.": "INR", "Rs": "INR", "INR": "INR",
    "£": "GBP", "GBP": "GBP",
    "$": "USD", "USD": "USD",
    "€": "EUR", "EUR": "EUR",
    "AED": "AED", "د.إ": "AED",
    "SAR": "SAR",
    "SGD": "SGD",
    "AUD": "AUD", "A$": "AUD",
    "CAD": "CAD", "C$": "CAD",
    "JPY": "JPY", "¥": "JPY",
    "CNY": "CNY",
}

# Regex that matches any currency amount like £125.00, Rs.50,843, $125.00, €50.00, INR 80,000
CURRENCY_AMT_RE = r"(?:₹|Rs\.?|INR|£|GBP|\$|USD|€|EUR|AED|د\.إ|SAR|SGD|A\$|AUD|C\$|CAD|¥|JPY|CNY)\s*[:\s]?\s*([\d,]+(?:\.\d+)?)"

HEADER_RE = re.compile(
    r"(?:Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday)"
    r",\s+(\d{1,2})\s+(\w+)\s+(?:(\d{4})\s+)?\xb7\s+(\d{2}):(\d{2})"
)
STANDALONE_TIME_RE = re.compile(r"^\s*\d{1,2}:\d{2}\s*$")

# Bank name detection
BANK_PATTERNS = [
    # Indian banks
    (re.compile(r"\bHDFC\s*Bank\b|\bHDFCBANK\b|\bHDFC\b", re.I), "HDFC Bank"),
    (re.compile(r"\bICICI\s*Bank\b|\bICICI\b", re.I), "ICICI Bank"),
    (re.compile(r"\bSBI\b|\bState Bank of India\b", re.I), "SBI"),
    (re.compile(r"\bAxis\s*Bank\b|\bAxis\b", re.I), "Axis Bank"),
    (re.compile(r"\bKotak\b", re.I), "Kotak Bank"),
    (re.compile(r"\bPNB\b|\bPunjab National Bank\b", re.I), "PNB"),
    (re.compile(r"\bBank of Baroda\b|\bBOB\b", re.I), "Bank of Baroda"),
    (re.compile(r"\bCanara Bank\b", re.I), "Canara Bank"),
    (re.compile(r"\bUnion Bank\b", re.I), "Union Bank"),
    (re.compile(r"\bStandard Chartered\b|\bSC Bank\b", re.I), "Standard Chartered"),
    (re.compile(r"\bIDBI Bank\b|\bIDBI\b", re.I), "IDBI Bank"),
    (re.compile(r"\bFederal Bank\b", re.I), "Federal Bank"),
    (re.compile(r"\bIndusInd Bank\b|\bIndusInd\b", re.I), "IndusInd Bank"),
    (re.compile(r"\bYes Bank\b", re.I), "Yes Bank"),
    (re.compile(r"\bRBL Bank\b|\bRatnakar Bank\b", re.I), "RBL Bank"),
    (re.compile(r"\bBandhan Bank\b", re.I), "Bandhan Bank"),
    (re.compile(r"\bIndian Bank\b", re.I), "Indian Bank"),
    (re.compile(r"\bIndian Overseas Bank\b|\bIOB\b", re.I), "IOB"),
    (re.compile(r"\bBank of India\b|\bBOI\b", re.I), "Bank of India"),
    (re.compile(r"\bCentral Bank\b", re.I), "Central Bank"),
    (re.compile(r"\bUCO Bank\b", re.I), "UCO Bank"),
    (re.compile(r"\bCity Union Bank\b", re.I), "City Union Bank"),
    (re.compile(r"\bDCB Bank\b", re.I), "DCB Bank"),
    (re.compile(r"\bKarur Vysya\b", re.I), "Karur Vysya Bank"),
    (re.compile(r"\bSouth Indian Bank\b", re.I), "South Indian Bank"),
    (re.compile(r"\bCSB Bank\b|\bCatholic Syrian\b", re.I), "CSB Bank"),
    (re.compile(r"\bKarnataka Bank\b", re.I), "Karnataka Bank"),
    (re.compile(r"\bTMB\b|\bTamilnad Mercantile\b", re.I), "TMB"),
    (re.compile(r"\bAU Small Finance\b", re.I), "AU SFB"),
    (re.compile(r"\bEquitas\b", re.I), "Equitas SFB"),
    (re.compile(r"\bUjjivan\b", re.I), "Ujjivan SFB"),
    (re.compile(r"\bFincare\b", re.I), "Fincare SFB"),
    (re.compile(r"\bJammu.*Kashmir Bank\b|\bJ&K Bank\b", re.I), "J&K Bank"),
    (re.compile(r"18002586465|1800\s*258\s*6465"), "Standard Chartered"),
    # UK banks
    (re.compile(r"\bLloyds\s*Bank\b", re.I), "Lloyds Bank"),
    (re.compile(r"\bBarclays\b", re.I), "Barclays"),
    (re.compile(r"\bHSBC\b", re.I), "HSBC"),
    (re.compile(r"\bNatWest\b", re.I), "NatWest"),
    (re.compile(r"\bSantander\b", re.I), "Santander"),
    (re.compile(r"\bHalifax\b", re.I), "Halifax"),
    (re.compile(r"\bNationwide\b", re.I), "Nationwide"),
    (re.compile(r"\bTSB\b", re.I), "TSB"),
    (re.compile(r"\bVirgin\s*Money\b", re.I), "Virgin Money"),
    (re.compile(r"\bFirst\s*Direct\b", re.I), "First Direct"),
    (re.compile(r"\bMonzo\b", re.I), "Monzo"),
    (re.compile(r"\bRevolut\b", re.I), "Revolut"),
    (re.compile(r"\bStarling\b", re.I), "Starling"),
    (re.compile(r"\bRoyal\s*Bank\s*of\s*Scotland\b|\bRBS\b", re.I), "RBS"),
    (re.compile(r"\bClydesdale\b", re.I), "Clydesdale Bank"),
    (re.compile(r"\bTesco\s*Bank\b", re.I), "Tesco Bank"),
    (re.compile(r"\bSainsburys\s*Bank\b", re.I), "Sainsburys Bank"),
    # US banks
    (re.compile(r"\bChase\b", re.I), "Chase"),
    (re.compile(r"\bBank\s*of\s*America\b|\bBofA\b", re.I), "Bank of America"),
    (re.compile(r"\bWells\s*Fargo\b", re.I), "Wells Fargo"),
    (re.compile(r"\bCitibank\b|\bCiti\b", re.I), "Citibank"),
    (re.compile(r"\bCapital\s*One\b", re.I), "Capital One"),
    (re.compile(r"\bUS\s*Bank\b", re.I), "US Bank"),
    (re.compile(r"\bPNC\s*Bank\b", re.I), "PNC Bank"),
    (re.compile(r"\bTD\s*Bank\b", re.I), "TD Bank"),
    (re.compile(r"\bTruist\b", re.I), "Truist"),
    (re.compile(r"\bDiscover\b", re.I), "Discover"),
    (re.compile(r"\bAmerican\s*Express\b|\bAmex\b", re.I), "American Express"),
    # Australian banks
    (re.compile(r"\bCommonwealth\s*Bank\b|\bCBA\b", re.I), "CommBank"),
    (re.compile(r"\bWestpac\b", re.I), "Westpac"),
    (re.compile(r"\bANZ\b", re.I), "ANZ"),
    (re.compile(r"\bNAB\b", re.I), "NAB"),
    # Canadian banks
    (re.compile(r"\bRBC\b|\bRoyal\s*Bank\b", re.I), "RBC"),
    (re.compile(r"\bTD\s*Canada\b|\bToronto\s*Dominion\b", re.I), "TD Canada"),
    (re.compile(r"\bScotiabank\b|\bBank\s*of\s*Nova\s*Scotia\b", re.I), "Scotiabank"),
    (re.compile(r"\bBMO\b|\bBank\s*of\s*Montreal\b", re.I), "BMO"),
    # Generic
    (re.compile(r"\bDeutsche\s*Bank\b", re.I), "Deutsche Bank"),
]


# ── Date parsing helpers ──────────────────────────────────────────────
def _safe_date(day_s, month_s, year_s, hour_s, minute_s, last_known_year):
    try:
        m = MONTH_MAP.get(month_s[:3].lower())
        if m is None:
            return ""
        yr = int(year_s) if year_s else last_known_year
        return datetime(yr, m, int(day_s), int(hour_s), int(minute_s)).strftime("%Y-%m-%d %H:%M")
    except Exception:
        return ""


def _parse_dd_mm(dd_mm, year_hint):
    try:
        p = dd_mm.split("-")
        return datetime(year_hint, int(p[1]), int(p[0])).strftime("%Y-%m-%d")
    except Exception:
        return ""


def _parse_ddmmyy(s):
    try:
        p = s.split("/")
        return datetime(2000 + int(p[2]), int(p[1]), int(p[0])).strftime("%Y-%m-%d")
    except Exception:
        return ""


def _parse_iso_datetime(s):
    try:
        return datetime.strptime(s, "%Y-%m-%d:%H:%M:%S").strftime("%Y-%m-%d %H:%M")
    except Exception:
        return ""


def _parse_compact_date(s):
    """Parse 'DD-MMM-YY' like '07-JUN-25'."""
    try:
        p = s.split("-")
        m = MONTH_MAP.get(p[1].lower())
        return datetime(2000 + int(p[2]), m, int(p[0])).strftime("%Y-%m-%d")
    except Exception:
        return ""


def _parse_month_dd(s):
    """Parse 'DD/MON/YYYY' like '07/JUN/2025'."""
    try:
        p = s.split("/")
        m = MONTH_MAP.get(p[1].lower())
        return datetime(int(p[2]), m, int(p[0])).strftime("%Y-%m-%d")
    except Exception:
        return ""


def _parse_compact_month(s):
    """Parse 'DDMonYY' like '06Jun26'."""
    try:
        m = MONTH_MAP.get(s[2:5].lower())
        d = int(s[:2])
        y = 2000 + int(s[5:7])
        return datetime(y, m, d).strftime("%Y-%m-%d")
    except Exception:
        return ""


def _parse_dd_mm_yy(s):
    """Parse DD/MM/YY or DD/MM/YYYY like '19/08/26' or '19/08/2026'."""
    try:
        p = s.split("/")
        day, mon = int(p[0]), int(p[1])
        yr = int(p[2])
        if yr < 100:
            yr += 2000
        return datetime(yr, mon, day).strftime("%Y-%m-%d")
    except Exception:
        return ""


def _clean_amount(raw):
    cleaned = raw.replace(",", "").strip()
    try:
        return float(cleaned)
    except ValueError:
        return None


def _detect_currency(text):
    """Detect currency from SMS text. Returns currency code."""
    for sym, code in CURRENCY_SYMBOLS.items():
        if sym in text:
            return code
    return DEFAULT_CURRENCY


def _detect_bank(text):
    for pattern, name in BANK_PATTERNS:
        if pattern.search(text):
            return name
    return "Unknown Bank"


def _is_non_transaction(text):
    """Check if SMS is a non-transaction notification to skip.
    IMPORTANT: Do NOT add patterns that appear inside valid transaction SMS.
    """
    # First: if the block contains a known transaction pattern, never skip it
    has_txn = re.search(
        r"Txn\s+Rs\.|Spent\s+Rs\.|IMPS\s+INR|Sent\s+Rs\.|"
        r"Payment Successful|debited\s+for\s+Rs\.|credited.*Rs\.|"
        r"deposited\s+in\b.*A/c|a/c\s+\w+\s+is\s+debited|"
        r"\d+\.\d+\s+spent\s+at\s+|\d+\.\d+\s+debit\s+|"
        r"\d+\.\d+\s+credit\s+|\d+\.\d+\s+payment\s+at\s+|"
        r"\d+\.\d+\s+withdrawal\s+|\d+\.\d+\s+transfer\s+to\s+|"
        r"\d+\.\d+\s+transfer\s+from\s+|"
        r"(?:£|\\\$|€)\d+\.\d+\s+(?:spent|debit|credit|payment|withdrawal|transfer)",
        text, re.I,
    )
    if has_txn:
        return False

    skip_patterns = [
        r"Available Bal(?:ance)?\s+in\b",
        r"Avail(?:ible)?\s+Bal(?:ance)?[\s:]+[£$€₹]",
        r"welcome kit",
        r"WELCOME KIT",
        r"Form Submitted",
        r"(?:^|\n)Amount Due\s*$",
        r"UPI Registration!",
        r"UPI LITE is successfully",
        r"Out for delivery",
        r"Debit Card will be delivered",
        r"maturity instruction",
        r"request.*case number.*is closed",
        r"has requested Rs.*frm u on Google Pay",
        r"has requested Rs.*frm u on PhonePe",
        r"has requested Rs.*frm u on Paytm",
        r"Thank you for considering.*Personal Loan",
        r"Credit Card Statement:\s*Total due",
        r"Email ID validation",
        r"Pending!\s*Email ID validation",
        r"kindly check your email inbox",
        r"FD\s+\w+\s+maturity instruction",
        r"Your FD No\b.*has been updated",
        r"Your .* Debit Card .* was issued on",
        r"will reach you soon",
        r"set PIN and controls",
        r"Linked.*to UPI",
        r"request to link your.*UPI",
        r"Messaging with.*is not available",
        r"not done by you",
        r"View this message on your phone",
        # UK-specific non-transactions
        r"Your\s+statement\s+is\s+ready",
        r"Your\s+card\s+has\s+been\s+blocked",
        r"New\s+direct\s+debit\s+set\s+up",
        r"Direct\s+debit\s+for\s+.*cancelled",
        r"Standing\s+order\s+has\s+been",
        r"Your\s+overdraft\s+(?:limit|facility)",
        r"Payment\s+(?:was\s+)?(?:received|processed)\s+but",
    ]
    for pat in skip_patterns:
        if re.search(pat, text, re.I):
            return True
    return False


# ── Transaction record ───────────────────────────────────────────────
class Txn:
    def __init__(self, date="", txn_type="", amount=None, category="",
                 merchant="", bank="", card="", ref="", detail="",
                 currency=None):
        self.date = date
        self.txn_type = txn_type
        self.amount = amount
        self.category = category
        self.merchant = merchant
        self.bank = bank
        self.card = card
        self.ref = ref
        self.detail = detail
        self.currency = currency or DEFAULT_CURRENCY


# ── Block splitting ──────────────────────────────────────────────────
def _split_into_blocks(lines):
    blocks = []
    current_date = ""
    current_lines = []
    last_known_year = datetime.now().year
    last_month_num = 0

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if STANDALONE_TIME_RE.match(stripped):
            current_lines.append(stripped)
            continue

        hdr = HEADER_RE.match(stripped)
        if hdr:
            if current_date or current_lines:
                blocks.append((current_date, current_lines))
            day_s, month_s, year_s, hh, mm = hdr.groups()
            cur_month_num = MONTH_MAP.get(month_s[:3].lower(), 0)
            if year_s:
                last_known_year = int(year_s)
                last_month_num = cur_month_num
            else:
                if last_month_num > 0 and cur_month_num < last_month_num:
                    last_known_year += 1
                last_month_num = cur_month_num
            current_date = _safe_date(day_s, month_s, year_s, hh, mm, last_known_year)
            current_lines = []
        else:
            current_lines.append(stripped)

    if current_date or current_lines:
        blocks.append((current_date, current_lines))
    return [(d, "\n".join(ls)) for d, ls in blocks]


# ── Parse one block ──────────────────────────────────────────────────
def _parse_block(header_date, body):
    if _is_non_transaction(body):
        return []

    txns = []
    # Split multi-transaction blocks on multiple transaction-start patterns
    split_pattern = (
        r"(?=Txn\s+Rs\.)"
        r"|(?=Your\s+a/c\s+\w+\s+is\s+debited)"
        r"|(?=ICICI\s+Bank\s+Account)"
        r"|(?=Dear\s+SBI)"
        r"|(?=your\s+A/c\s+\w+-?credited)"
        r"|(?=Sent\s+Rs\.)"
        r"|(?=IMPS\s+INR)"
        r"|(?=Spent\s+Rs\.)"
        r"|(?=HDFC Bank:Rs\.)"
        r"|(?=Payment Successful!)"
        r"|(?=UPDATE:\s*INR)"
        r"|(?=Update!\s*INR)"
        r"|(?=FD\s+\w+\s+of\s+INR)"
        r"|(?=Mandate\s+Set)"
        r"|(?=Paid\s+Rs\.)"
        r"|(?=\d+\.\d+\s+spent\s+at\s+)"
        r"|(?=\w+\s*Bank:\s*[£$€₹])"
        r"|(?=\w+\s*Bank:\s*\d)"
        r"|(?=Your\s+\w+\s+Bank\b)"
        r"|(?=Lloyds\s*Bank:)"
    )
    sub_parts = re.split(split_pattern, body)
    for part in sub_parts:
        part = part.strip()
        if not part:
            continue
        txn = _parse_single_text(header_date, part)
        if txn is not None:
            txns.append(txn)
    return txns


def _parse_single_text(header_date, text):
    """Match a single transaction from text. Tries all bank patterns."""
    bank = _detect_bank(text)
    currency = _detect_currency(text)

    # ── Lloyds Bank / UK generic: "£125.00 spent at Tesco on 19/08/26 14:05" ──
    m = re.search(
        r"([£$€₹])\s*([\d,]+(?:\.\d+)?)\s+spent\s+at\s+(.+?)\s+on\s+"
        r"(\d{1,2}/\d{1,2}/\d{2,4})\s+(\d{2}:\d{2})",
        text, re.S,
    )
    if m:
        sym, amt_s, merchant, date_s, time_s = m.groups()
        cur = CURRENCY_SYMBOLS.get(sym, currency)
        return Txn(
            date=_parse_dd_mm_yy(date_s) + " " + time_s if _parse_dd_mm_yy(date_s) else header_date,
            txn_type="Debit", amount=_clean_amount(amt_s),
            category="CC Spent", merchant=merchant.strip(),
            bank=bank, currency=cur,
        )

    # ── Lloyds/UK generic: "£125.00 debit at Tesco" ──
    m = re.search(
        r"([£$€₹])\s*([\d,]+(?:\.\d+)?)\s+debit\s+(?:at|from)\s+(.+?)(?:\s+on\s+(\d{1,2}/\d{1,2}/\d{2,4})\s+(\d{2}:\d{2}))?",
        text, re.S,
    )
    if m:
        sym, amt_s, merchant, date_s, time_s = m.groups()
        cur = CURRENCY_SYMBOLS.get(sym, currency)
        dt = (header_date or "")
        if date_s:
            parsed = _parse_dd_mm_yy(date_s)
            if parsed:
                dt = parsed + (" " + time_s if time_s else "")
        return Txn(
            date=dt, txn_type="Debit", amount=_clean_amount(amt_s),
            category="Debit Card", merchant=merchant.strip(),
            bank=bank, currency=cur,
        )

    # ── Lloyds/UK generic: "£125.00 credit from Employer" ──
    m = re.search(
        r"([£$€₹])\s*([\d,]+(?:\.\d+)?)\s+credit\s+from\s+(.+?)(?:\s+on\s+(\d{1,2}/\d{1,2}/\d{2,4})\s+(\d{2}:\d{2}))?",
        text, re.S,
    )
    if m:
        sym, amt_s, merchant, date_s, time_s = m.groups()
        cur = CURRENCY_SYMBOLS.get(sym, currency)
        dt = (header_date or "")
        if date_s:
            parsed = _parse_dd_mm_yy(date_s)
            if parsed:
                dt = parsed + (" " + time_s if time_s else "")
        return Txn(
            date=dt, txn_type="Credit", amount=_clean_amount(amt_s),
            category="Credit", merchant=merchant.strip(),
            bank=bank, currency=cur,
        )

    # ── Lloyds/UK generic: "£125.00 payment to Tesco" ──
    m = re.search(
        r"([£$€₹])\s*([\d,]+(?:\.\d+)?)\s+payment\s+to\s+(.+?)(?:\s+on\s+(\d{1,2}/\d{1,2}/\d{2,4})\s+(\d{2}:\d{2}))?",
        text, re.S,
    )
    if m:
        sym, amt_s, merchant, date_s, time_s = m.groups()
        cur = CURRENCY_SYMBOLS.get(sym, currency)
        dt = (header_date or "")
        if date_s:
            parsed = _parse_dd_mm_yy(date_s)
            if parsed:
                dt = parsed + (" " + time_s if time_s else "")
        return Txn(
            date=dt, txn_type="Debit", amount=_clean_amount(amt_s),
            category="Payment", merchant=merchant.strip(),
            bank=bank, currency=cur,
        )

    # ── Lloyds/UK generic: "£125.00 transfer to Name" ──
    m = re.search(
        r"([£$€₹])\s*([\d,]+(?:\.\d+)?)\s+transfer\s+to\s+(.+?)(?:\s+on\s+(\d{1,2}/\d{1,2}/\d{2,4})\s+(\d{2}:\d{2}))?",
        text, re.S,
    )
    if m:
        sym, amt_s, merchant, date_s, time_s = m.groups()
        cur = CURRENCY_SYMBOLS.get(sym, currency)
        dt = (header_date or "")
        if date_s:
            parsed = _parse_dd_mm_yy(date_s)
            if parsed:
                dt = parsed + (" " + time_s if time_s else "")
        return Txn(
            date=dt, txn_type="Debit", amount=_clean_amount(amt_s),
            category="Transfer", merchant=merchant.strip(),
            bank=bank, currency=cur,
        )

    # ── Lloyds/UK generic: "£125.00 transfer from Name" ──
    m = re.search(
        r"([£$€₹])\s*([\d,]+(?:\.\d+)?)\s+transfer\s+from\s+(.+?)(?:\s+on\s+(\d{1,2}/\d{1,2}/\d{2,4})\s+(\d{2}:\d{2}))?",
        text, re.S,
    )
    if m:
        sym, amt_s, merchant, date_s, time_s = m.groups()
        cur = CURRENCY_SYMBOLS.get(sym, currency)
        dt = (header_date or "")
        if date_s:
            parsed = _parse_dd_mm_yy(date_s)
            if parsed:
                dt = parsed + (" " + time_s if time_s else "")
        return Txn(
            date=dt, txn_type="Credit", amount=_clean_amount(amt_s),
            category="Transfer", merchant=merchant.strip(),
            bank=bank, currency=cur,
        )

    # ── Lloyds/UK generic: "£125.00 withdrawal at ATM" ──
    m = re.search(
        r"([£$€₹])\s*([\d,]+(?:\.\d+)?)\s+withdrawal\s+(?:at|from)\s+(.+?)(?:\s+on\s+(\d{1,2}/\d{1,2}/\d{2,4})\s+(\d{2}:\d{2}))?",
        text, re.S,
    )
    if m:
        sym, amt_s, merchant, date_s, time_s = m.groups()
        cur = CURRENCY_SYMBOLS.get(sym, currency)
        dt = (header_date or "")
        if date_s:
            parsed = _parse_dd_mm_yy(date_s)
            if parsed:
                dt = parsed + (" " + time_s if time_s else "")
        return Txn(
            date=dt, txn_type="Debit", amount=_clean_amount(amt_s),
            category="ATM Withdrawal", merchant=merchant.strip(),
            bank=bank, currency=cur,
        )

    # ── HDFC: CC UPI Txn ────────────────────────────────────────
    m = re.search(
        r"Txn\s+Rs\.([\d,]+(?:\.\d+)?)"
        r".*?(?:On|on)\s+\w+\s+Bank\s+Card\s+(\d+)"
        r".*?At\s+(.+?)\s+by\s+UPI\s+(\d+)"
        r".*?On\s+(\d{2})-(\d{2})",
        text, re.S,
    )
    if m:
        return Txn(
            date=header_date, txn_type="Debit", amount=_clean_amount(m.group(1)),
            category="CC UPI", merchant=m.group(3).strip(),
            bank=bank, card=m.group(2), ref=m.group(4), currency=currency,
        )

    # ── HDFC: CC Spent ──────────────────────────────────────────
    m = re.search(
        r"Spent\s+Rs\.([\d,]+(?:\.\d+)?)"
        r".*?Bank\s+Card\s+(\d+)"
        r".*?At\s+(.+?)\s+On\s+(\d{4}-\d{2}-\d{2}:\d{2}:\d{2}:\d{2})",
        text, re.S,
    )
    if m:
        return Txn(
            date=_parse_iso_datetime(m.group(4)) or header_date,
            txn_type="Debit", amount=_clean_amount(m.group(1)),
            category="CC Spent", merchant=m.group(3).strip(),
            bank=bank, card=m.group(2), currency=currency,
        )

    # ── ICICI: CC Spent ─────────────────────────────────────────
    m = re.search(
        r"INR\s+([\d,]+(?:\.\d+)?)\s+spent\s+using\s+\w+\s+Card\s+(\w+)"
        r"\s+on\s+(\d{2})-(\w{3})-(\d{2})"
        r"\s+on\s+(.+?)(?:\.|\s+Avl|\s+If not)",
        text, re.S,
    )
    if m:
        mon = MONTH_MAP.get(m.group(4).lower())
        if mon:
            try:
                dt = datetime(2000 + int(m.group(5)), mon, int(m.group(3))).strftime("%Y-%m-%d")
            except Exception:
                dt = header_date
        else:
            dt = header_date
        return Txn(
            date=dt, txn_type="Debit", amount=_clean_amount(m.group(1)),
            category="CC Spent", merchant=m.group(6).strip(),
            bank=bank, card=m.group(2), currency=currency,
        )

    # ── HDFC: IMPS sent ─────────────────────────────────────────
    m = re.search(
        r"IMPS\s+INR\s+([\d,]+(?:\.\d+)?)"
        r".*?sent from\s+\w+\s+Bank\s+A/c\s+(\w+)"
        r".*?on\s+(\d{2})-(\d{2})-(\d{2})"
        r".*?To A/c\s+(\S+)"
        r".*?Ref-(\S+)",
        text, re.S,
    )
    if m:
        return Txn(
            date=header_date, txn_type="Debit", amount=_clean_amount(m.group(1)),
            category="IMPS", merchant=f"To {m.group(6)}",
            bank=bank, ref=m.group(7), detail=f"From {m.group(2)}", currency=currency,
        )

    # ── HDFC: UPI Sent ──────────────────────────────────────────
    m = re.search(
        r"Sent\s+Rs\.([\d,]+(?:\.\d+)?)"
        r".*?From\s+\w+\s+Bank\s+A/[Cc]\s+(\S+)"
        r".*?To\s+(.+?)\s+On\s+(\d{2}/\d{2}/\d{2})"
        r".*?Ref\s+(\S+)",
        text, re.S,
    )
    if m:
        return Txn(
            date=header_date, txn_type="Debit", amount=_clean_amount(m.group(1)),
            category="UPI Sent", merchant=m.group(3).strip(),
            bank=bank, ref=m.group(5), detail=f"From {m.group(2)}", currency=currency,
        )

    # ── Generic: a/c debited UPI (SC, SBI, all banks) ───────────
    m = re.search(
        r"(?:Your\s+)?a/c\s+(\w+)\s+is\s+debited\s+for\s+Rs\.?\s*([\d,]+(?:\.\d+)?)"
        r"\s+on\s+(\d{2}-\d{2}-\d{4})\s+(\d{2}:\d{2})"
        r".*?credited\s+to\s+a/c\s+(\S+)"
        r".*?(?:UPI\s+Ref\s+no\.?\s*(\d+))?",
        text, re.S,
    )
    if m:
        try:
            dt = datetime.strptime(m.group(3) + " " + m.group(4), "%d-%m-%Y %H:%M").strftime("%Y-%m-%d %H:%M")
        except Exception:
            dt = header_date
        return Txn(
            date=dt, txn_type="Debit", amount=_clean_amount(m.group(2)),
            category="UPI Debit", merchant=f"To {m.group(5)}",
            bank=bank, card=m.group(1), ref=m.group(6) or "", currency=currency,
        )

    # ── HDFC: Account Debit (UPDATE: INR ... debited) ───────────
    m = re.search(
        r"UPDATE:\s*INR\s+([\d,]+(?:\.\d+)?)"
        r"\s+debited from\s+\w+\s+Bank\s+(\w+)"
        r".*?on\s+(\d{2}-\w{3}-\d{2})"
        r".*?Info:\s*(.+?)(?:\.|\s+Avl|\s*$)",
        text, re.S,
    )
    if m:
        return Txn(
            date=_parse_compact_date(m.group(3)), txn_type="Debit",
            amount=_clean_amount(m.group(1)),
            category="Account Debit", merchant=m.group(4).strip().rstrip("."),
            bank=bank, detail=f"A/c {m.group(2)}", currency=currency,
        )

    # ── HDFC: Payment Successful (NetBanking) ───────────────────
    m = re.search(
        r"Payment Successful!\s*Rs\.?\s*([\d,]+(?:\.\d+)?)"
        r".*?from A/c\s+(\S+)"
        r".*?to\s+(\S+)",
        text, re.S,
    )
    if m:
        return Txn(
            date=header_date, txn_type="Debit", amount=_clean_amount(m.group(1)),
            category="NetBanking", merchant=f"To {m.group(3)}",
            bank=bank, detail=f"From {m.group(2)}", currency=currency,
        )

    # ── HDFC: Debit from account (HDFC Bank:Rs. ... debited) ────
    m = re.search(
        r"HDFC Bank:Rs\.?\s*([\d,]+(?:\.\d+)?)"
        r"\s+debited from a/c\s+(\S+)"
        r".*?to a/c\s+(\S+)"
        r".*?UPI Ref No\.?\s*(\d+)",
        text, re.S,
    )
    if m:
        return Txn(
            date=header_date, txn_type="Debit", amount=_clean_amount(m.group(1)),
            category="UPI Debit", merchant=f"To {m.group(3)}",
            bank="HDFC Bank", card=m.group(2), ref=m.group(4), currency=currency,
        )

    # ── HDFC: Bill Payment (Paid Rs. For: ...) ──────────────────
    m = re.search(
        r"Paid\s+Rs\.?\s*([\d,]+(?:\.\d+)?)"
        r".*?For:\s*(.+?)\s+From\s+\w+\s+Bank\s+A/c\s+(\S+)",
        text, re.S,
    )
    if m:
        return Txn(
            date=header_date, txn_type="Debit", amount=_clean_amount(m.group(1)),
            category="Bill Payment", merchant=m.group(2).strip(),
            bank=bank, detail=f"From {m.group(3)}", currency=currency,
        )

    # ── HDFC: Mandate Set ───────────────────────────────────────
    m = re.search(
        r"Mandate\s+Set\s+Rs\.([\d,]+(?:\.\d+)?)"
        r".*?For\s+(.+?)\s+From\s+\w+\s+Bank\s+[Aa]/[Cc]\s+(\S+)",
        text, re.S,
    )
    if m:
        return Txn(
            date=header_date, txn_type="Debit", amount=_clean_amount(m.group(1)),
            category="Mandate", merchant=m.group(2).strip(),
            bank=bank, detail=f"From {m.group(3)}", currency=currency,
        )

    # ── HDFC: FD Sweep-In Credit ────────────────────────────────
    m = re.search(
        r"Update!\s*INR\s+([\d,]+(?:\.\d+)?)"
        r"\s+deposited in\s+\w+\s+Bank\s+A/c\s+(\w+)"
        r".*?on\s+(\d{2}-\w{3}-\d{2})"
        r".*?SWEEP-IN CREDIT\s*-\s*(\S+)",
        text, re.S,
    )
    if m:
        return Txn(
            date=_parse_compact_date(m.group(3)), txn_type="Credit",
            amount=_clean_amount(m.group(1)),
            category="FD Sweep-In", merchant="Sweep-In Credit",
            bank=bank, detail=f"A/c {m.group(2)} Ref {m.group(4)}", currency=currency,
        )

    # ── HDFC: FD Liquidation ────────────────────────────────────
    m = re.search(
        r"FD\s+(\w+)\s+of\s+INR\s+([\d,]+(?:\.\d+)?)"
        r"\s+will be liquidated for INR\s+([\d,]+(?:\.\d+)?)"
        r".*?from A/c\s+(\w+)",
        text, re.S,
    )
    if m:
        return Txn(
            date=header_date, txn_type="Debit",
            amount=_clean_amount(m.group(3)),
            category="FD Liquidation",
            merchant=f"FD {m.group(1)} (INR {_clean_amount(m.group(2))})",
            bank=bank, detail=f"From A/c {m.group(4)}", currency=currency,
        )

    # ── ICICI: Account Credit ────────────────────────────────────
    m = re.search(
        r"ICICI Bank Account\s+(\w+)\s+credited\s*:\s*Rs\.?\s*([\d,]+(?:\.\d+)?)"
        r"\s+on\s+(\d{2})-(\w{3})-(\d{2})"
        r"(?:\.?\s*Info\s*(.+?))?(?:\.?\s*Available|\s*$)",
        text, re.S,
    )
    if m:
        mon = MONTH_MAP.get(m.group(4).lower())
        if mon:
            try:
                dt = datetime(2000 + int(m.group(5)), mon, int(m.group(3))).strftime("%Y-%m-%d")
            except Exception:
                dt = header_date
        else:
            dt = header_date
        info = (m.group(6) or "").strip().rstrip(".")
        return Txn(
            date=dt, txn_type="Credit", amount=_clean_amount(m.group(2)),
            category="Credit", merchant=info or "Credit",
            bank="ICICI Bank", card=m.group(1), currency=currency,
        )

    # ── ICICI: IMPS Credit ──────────────────────────────────────
    m = re.search(
        r"ICICI Bank Account\s+(\w+)\s+is\s+credited\s+with\s+Rs\.?\s*([\d,]+(?:\.\d+)?)"
        r"\s+on\s+(\d{2})-(\w{3})-(\d{2})"
        r"\s+by\s+(.+?)\.?\s+IMPS\s+Ref\.?\s*no\.?\s*(\d+)",
        text, re.S,
    )
    if m:
        mon = MONTH_MAP.get(m.group(4).lower())
        if mon:
            try:
                dt = datetime(2000 + int(m.group(5)), mon, int(m.group(3))).strftime("%Y-%m-%d")
            except Exception:
                dt = header_date
        else:
            dt = header_date
        return Txn(
            date=dt, txn_type="Credit", amount=_clean_amount(m.group(2)),
            category="IMPS Credit", merchant=m.group(6).strip(),
            bank="ICICI Bank", card=m.group(1), ref=m.group(7), currency=currency,
        )

    # ── SBI: Account Credit ─────────────────────────────────────
    m = re.search(
        r"Dear\s+SBI\s+User.*?A/c\s+(\w+)-credited\s+by\s+Rs\.?([\d,]+(?:\.\d+)?)"
        r"\s+on\s+(\d{2})(\w{3})(\d{2,4})"
        r".*?transfer\s+from\s+(.+?)\s+Ref\s+No\s+(\d+)",
        text, re.S,
    )
    if m:
        yr_str = m.group(5)
        yr = int(yr_str) if len(yr_str) == 4 else 2000 + int(yr_str)
        mon = MONTH_MAP.get(m.group(4).lower())
        if mon:
            try:
                dt = datetime(yr, mon, int(m.group(3))).strftime("%Y-%m-%d")
            except Exception:
                dt = header_date
        else:
            dt = header_date
        return Txn(
            date=dt, txn_type="Credit", amount=_clean_amount(m.group(2)),
            category="Credit", merchant=m.group(6).strip(),
            bank="SBI", card=m.group(1), ref=m.group(7), currency=currency,
        )

    # ── SBI: Generic Credit ─────────────────────────────────────
    m = re.search(
        r"your\s+A/c\s+(\w+)-?credited\s+by\s+Rs\.?([\d,]+(?:\.\d+)?)"
        r"\s+on\s+(\d{2})(\w{3})(\d{2,4})"
        r".*?transfer\s+from\s+(.+?)\s+Ref\s+No\s+(\d+)",
        text, re.S,
    )
    if m:
        yr_str = m.group(5)
        yr = int(yr_str) if len(yr_str) == 4 else 2000 + int(yr_str)
        mon = MONTH_MAP.get(m.group(4).lower())
        if mon:
            try:
                dt = datetime(yr, mon, int(m.group(3))).strftime("%Y-%m-%d")
            except Exception:
                dt = header_date
        else:
            dt = header_date
        return Txn(
            date=dt, txn_type="Credit", amount=_clean_amount(m.group(2)),
            category="Credit", merchant=m.group(6).strip(),
            bank=_detect_bank(text), card=m.group(1), ref=m.group(7), currency=currency,
        )

    # ── HDFC: Card Payment Credit ───────────────────────────────
    m = re.search(
        r"Online Payment of Rs\.([\d,]+(?:\.\d+)?)"
        r".*?card ending\s+(\d+)"
        r".*?On\s+(\d{2}/\w+/\d{4})",
        text, re.S,
    )
    if m:
        return Txn(
            date=_parse_month_dd(m.group(3)) or header_date,
            txn_type="Credit", amount=_clean_amount(m.group(1)),
            category="CC Payment", merchant="Card Payment",
            bank=bank, card=m.group(2), currency=currency,
        )

    m = re.search(
        r"Payment of Rs\s+([\d,]+(?:\.\d+)?)\s+was credited"
        r".*?card ending\s+(\d+)"
        r".*?On\s+(\d{2}/\w+/\d{4})",
        text, re.S,
    )
    if m:
        return Txn(
            date=_parse_month_dd(m.group(3)) or header_date,
            txn_type="Credit", amount=_clean_amount(m.group(1)),
            category="CC Payment", merchant="Card Payment",
            bank=bank, card=m.group(2), currency=currency,
        )

    # ── Generic fallback: any currency "spent at" / "debit" / "credit" ──
    # Catches patterns not matched above: e.g. "$50.00 spent at Walmart"
    m = re.search(
        r"([£$€₹])\s*([\d,]+(?:\.\d+)?)\s+(spent\s+at|debit|credit|payment|withdrawal|transfer)\s+"
        r"(?:(?:to|at|from)\s+)?(.+?)(?:\s+on\s+(\d{1,2}/\d{1,2}/\d{2,4})\s+(\d{2}:\d{2}))?",
        text, re.S,
    )
    if m:
        sym, amt_s, action, target, date_s, time_s = m.groups()
        cur = CURRENCY_SYMBOLS.get(sym, currency)
        dt = (header_date or "")
        if date_s:
            parsed = _parse_dd_mm_yy(date_s)
            if parsed:
                dt = parsed + (" " + time_s if time_s else "")
        is_credit = "credit" in action.lower() or "from" in action.lower()
        return Txn(
            date=dt,
            txn_type="Credit" if is_credit else "Debit",
            amount=_clean_amount(amt_s),
            category=action.strip().title(),
            merchant=target.strip(),
            bank=bank, currency=cur,
        )

    return None


# ── Main parser ──────────────────────────────────────────────────────
def parse_sms_file(path):
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    blocks = _split_into_blocks(lines)
    txns = []
    for header_date, body in blocks:
        block_txns = _parse_block(header_date, body)
        txns.extend(block_txns)
    return txns


def parse_sms_text(text):
    """Parse SMS from raw text (for Streamlit file upload / paste)."""
    lines = text.splitlines()
    blocks = _split_into_blocks(lines)
    txns = []
    for header_date, body in blocks:
        block_txns = _parse_block(header_date, body)
        txns.extend(block_txns)
    return txns


# ── Excel export ──────────────────────────────────────────────────────
HEADERS = ["Date", "Type", "Amount", "Currency", "Category", "Merchant / Detail",
           "Bank", "Card / Account", "Ref / Note"]

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
CREDIT_FONT = Font(color="006100")
DEBIT_FONT  = Font(color="9C0006")
CREDIT_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
DEBIT_FILL  = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
NUM_FMT = '#,##0.00'

# Currency symbols for display
CURRENCY_DISPLAY = {
    "INR": "₹", "GBP": "£", "USD": "$", "EUR": "€",
    "AED": "AED ", "SAR": "SAR ", "SGD": "S$",
    "AUD": "A$", "CAD": "C$", "JPY": "¥", "CNY": "¥",
}


def export_to_excel(txns, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"

    for col, hdr in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=hdr)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER

    for idx, t in enumerate(txns, 2):
        cur_sym = CURRENCY_DISPLAY.get(t.currency or DEFAULT_CURRENCY, "")
        row_data = [t.date, t.txn_type, t.amount, t.currency or DEFAULT_CURRENCY,
                    t.category, t.merchant, t.bank, t.card or t.detail, t.ref]
        fill = CREDIT_FILL if t.txn_type == "Credit" else DEBIT_FILL
        font = CREDIT_FONT if t.txn_type == "Credit" else DEBIT_FONT
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=idx, column=col, value=val)
            cell.border = THIN_BORDER
            cell.fill = fill
            cell.font = font
            if col == 3 and val is not None:
                cell.number_format = NUM_FMT
                cell.alignment = Alignment(horizontal="right")

    widths = [19, 8, 14, 10, 16, 38, 18, 28, 20]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    total_credit = sum(t.amount for t in txns if t.txn_type == "Credit" and t.amount)
    total_debit  = sum(t.amount for t in txns if t.txn_type == "Debit"  and t.amount)

    # Group by currency
    by_currency = {}
    for t in txns:
        cur = t.currency or DEFAULT_CURRENCY
        if cur not in by_currency:
            by_currency[cur] = {"credit": 0, "debit": 0, "count": 0}
        if t.txn_type == "Credit" and t.amount:
            by_currency[cur]["credit"] += t.amount
        elif t.txn_type == "Debit" and t.amount:
            by_currency[cur]["debit"] += t.amount
        by_currency[cur]["count"] += 1

    summary = [
        ("Total Transactions", len(txns)),
    ]
    for cur, totals in sorted(by_currency.items()):
        sym = CURRENCY_DISPLAY.get(cur, cur + " ")
        summary.append((f"Total Credits ({cur})", f"{sym}{totals['credit']:,.2f}"))
        summary.append((f"Total Debits ({cur})",  f"{sym}{totals['debit']:,.2f}"))
        summary.append((f"Net ({cur})",           f"{sym}{totals['credit'] - totals['debit']:,.2f}"))
        summary.append(("", ""))

    summary.append(("Category Breakdown (Debits)", ""))
    cat_debits = {}
    for t in txns:
        if t.txn_type == "Debit" and t.amount:
            cat_debits[t.category] = cat_debits.get(t.category, 0) + t.amount
    for cat in sorted(cat_debits, key=cat_debits.get, reverse=True):
        summary.append((cat, f"{cat_debits[cat]:,.2f}"))

    summary.append(("", ""))
    summary.append(("Banks Detected", ""))
    bank_counts = {}
    for t in txns:
        bank_counts[t.bank] = bank_counts.get(t.bank, 0) + 1
    for b, cnt in sorted(bank_counts.items(), key=lambda x: -x[1]):
        summary.append((b, f"{cnt} txns"))

    for r, (label, val) in enumerate(summary, 1):
        c1 = ws2.cell(row=r, column=1, value=label)
        c2 = ws2.cell(row=r, column=2, value=val)
        c1.font = Font(bold=True, size=11)
        c2.alignment = Alignment(horizontal="right")
        c1.border = THIN_BORDER
        c2.border = THIN_BORDER
    ws2.column_dimensions["A"].width = 32
    ws2.column_dimensions["B"].width = 20

    wb.save(out_path)
    return out_path


# ── CLI entry point ──────────────────────────────────────────────────
if __name__ == "__main__":
    default_in  = r"D:\SMSTransactionAnalysis\DailyExpenseAnalyser\dailytran.txt"
    default_out = r"D:\SMSTransactionAnalysis\DailyExpenseAnalyser\transactions.xlsx"

    input_file  = sys.argv[1] if len(sys.argv) > 1 else default_in
    output_file = sys.argv[2] if len(sys.argv) > 2 else default_out

    print(f"Reading: {input_file}")
    print(f"Default currency: {DEFAULT_CURRENCY}")
    txns = parse_sms_file(input_file)
    print(f"Parsed {len(txns)} transactions")

    if not txns:
        print("No transactions found.")
        sys.exit(1)

    out = export_to_excel(txns, output_file)
    print(f"Saved to: {out}")

    # Group totals by currency
    by_cur = {}
    for t in txns:
        cur = t.currency or DEFAULT_CURRENCY
        if cur not in by_cur:
            by_cur[cur] = {"credit": 0, "debit": 0}
        if t.txn_type == "Credit" and t.amount:
            by_cur[cur]["credit"] += t.amount
        elif t.txn_type == "Debit" and t.amount:
            by_cur[cur]["debit"] += t.amount

    for cur, totals in sorted(by_cur.items()):
        sym = CURRENCY_DISPLAY.get(cur, cur + " ")
        print(f"\n[{cur}]")
        print(f"  Credits: {sym}{totals['credit']:,.2f}")
        print(f"  Debits:  {sym}{totals['debit']:,.2f}")
        print(f"  Net:     {sym}{totals['credit'] - totals['debit']:,.2f}")
